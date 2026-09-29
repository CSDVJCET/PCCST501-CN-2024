"""
build_all_topic_ppts.py
Generates professional, modern, beautifully styled PowerPoint presentations (.pptx)
for PCCST501 Computer Networks (KTU 2024 Scheme) covering all topics across Modules 1 to 4.
Prepared by: Prof. Anju Markose | Assistant Professor, Dept. of CSE | VJCET
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# --- Refined Theme Palette ---
COLOR_PRIMARY = RGBColor(15, 30, 74)       # Deep Navy Blue
COLOR_SECONDARY = RGBColor(14, 116, 144)   # Deep Teal / Cyan
COLOR_ACCENT = RGBColor(217, 119, 6)       # Warm Amber / Gold
COLOR_BG_DARK = RGBColor(15, 23, 42)       # Slate 900
COLOR_TEXT_LIGHT = RGBColor(248, 250, 252) # Slate 50
COLOR_TEXT_DARK = RGBColor(30, 41, 59)     # Slate 800
COLOR_TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
COLOR_CARD_BG = RGBColor(241, 245, 249)    # Slate 100
COLOR_CARD_BORDER = RGBColor(203, 213, 225)# Slate 300
COLOR_CARD_ACCENT_BG = RGBColor(238, 242, 255) # Indigo 50

AUTHOR_TAG = "Prof. Anju Markose | Assistant Professor, Dept. of CSE | VJCET"
COURSE_TAG = "PCCST501 Computer Networks | KTU 2024 Scheme"

def create_title_slide(prs, topic_code, title, subtitle, module_name):
    """Creates a sleek, dark-themed title slide with gold accent bar and author details."""
    slide_layout = prs.slide_layouts[6] # blank layout
    slide = prs.slides.add_slide(slide_layout)
    
    # Dark Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG_DARK
    bg.line.fill.background()
    
    # Left Vertical Gold Accent Bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.4), Inches(0.18), Inches(4.7))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_ACCENT
    bar.line.fill.background()
    
    # Text Container
    tx_box = slide.shapes.add_textbox(Inches(1.4), Inches(1.3), Inches(11.0), Inches(4.9))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    # Module & Topic Tag
    p0 = tf.paragraphs[0]
    p0.text = f"{COURSE_TAG} • {module_name.upper()} • {topic_code}"
    p0.font.bold = True
    p0.font.size = Pt(13)
    p0.font.color.rgb = COLOR_ACCENT
    p0.space_after = Pt(12)
    
    # Main Title
    p1 = tf.add_paragraph()
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(32)
    p1.font.color.rgb = COLOR_TEXT_LIGHT
    p1.space_after = Pt(14)
    
    # Subtitle / Overview
    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.size = Pt(17)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.space_after = Pt(28)
    
    # Author & Institution Signature
    p3 = tf.add_paragraph()
    p3.text = f"Prepared by: {AUTHOR_TAG}"
    p3.font.bold = True
    p3.font.size = Pt(13)
    p3.font.color.rgb = RGBColor(226, 232, 240)

    p4 = tf.add_paragraph()
    p4.text = "Viswajyothi College of Engineering and Technology (VJCET), Vazhakulam"
    p4.font.size = Pt(11)
    p4.font.color.rgb = RGBColor(148, 163, 184)
    
    return slide

def create_content_slide(prs, title, topic_tag, points, card_header=None, card_content=None):
    """Creates a structured content slide with top primary banner and optional right takeaway card."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Top Banner
    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_PRIMARY
    banner.line.fill.background()
    
    # Header Accent Line
    banner_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.06), Inches(13.333), Inches(0.04))
    banner_line.fill.solid()
    banner_line.fill.fore_color.rgb = COLOR_ACCENT
    banner_line.line.fill.background()
    
    # Slide Title
    tx_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.18), Inches(9.2), Inches(0.75))
    tf_t = tx_title.text_frame
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.bold = True
    p_t.font.size = Pt(23)
    p_t.font.color.rgb = COLOR_TEXT_LIGHT
    
    # Tag Text (Top Right)
    tx_tag = slide.shapes.add_textbox(Inches(9.0), Inches(0.18), Inches(3.6), Inches(0.75))
    tf_tag = tx_tag.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = topic_tag
    p_tag.alignment = PP_ALIGN.RIGHT
    p_tag.font.bold = True
    p_tag.font.size = Pt(13)
    p_tag.font.color.rgb = COLOR_ACCENT
    
    # Left Content Container
    left_w = Inches(7.5) if card_header else Inches(11.8)
    tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), left_w, Inches(5.4))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    for i, pt in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(pt, tuple): # (Heading, Details)
            p.text = f"• {pt[0]}"
            p.font.bold = True
            p.font.size = Pt(15)
            p.font.color.rgb = COLOR_PRIMARY
            p.space_after = Pt(2)
            
            p_sub = tf.add_paragraph()
            p_sub.text = f"   {pt[1]}"
            p_sub.font.size = Pt(13.5)
            p_sub.font.color.rgb = COLOR_TEXT_DARK
            p_sub.space_after = Pt(8)
        else:
            p.text = f"• {pt}"
            p.font.size = Pt(14)
            p.font.color.rgb = COLOR_TEXT_DARK
            p.space_after = Pt(8)
            
    # Right Highlight Card
    if card_header and card_content:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.7), Inches(1.4), Inches(3.9), Inches(5.3))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_SECONDARY
        card.line.width = Pt(1.5)
        
        tx_card = slide.shapes.add_textbox(Inches(8.9), Inches(1.55), Inches(3.5), Inches(5.0))
        tf_c = tx_card.text_frame
        tf_c.word_wrap = True
        
        p_ch = tf_c.paragraphs[0]
        p_ch.text = card_header
        p_ch.font.bold = True
        p_ch.font.size = Pt(15)
        p_ch.font.color.rgb = COLOR_SECONDARY
        p_ch.space_after = Pt(10)
        
        for c_pt in card_content:
            p_c = tf_c.add_paragraph()
            p_c.text = f"→ {c_pt}"
            p_c.font.size = Pt(12.5)
            p_c.font.color.rgb = COLOR_TEXT_DARK
            p_c.space_after = Pt(7)
            
    # Slide Footer
    footer = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.8), Inches(0.4))
    tf_f = footer.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = f"{COURSE_TAG} | {AUTHOR_TAG}"
    p_f.font.size = Pt(9.5)
    p_f.font.color.rgb = COLOR_TEXT_MUTED
    
    return slide

def create_comparison_slide(prs, title, topic_tag, col1_title, col1_items, col2_title, col2_items, col3_title=None, col3_items=None):
    """Creates a multi-column comparison slide."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Top Banner
    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_PRIMARY
    banner.line.fill.background()
    
    banner_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.06), Inches(13.333), Inches(0.04))
    banner_line.fill.solid()
    banner_line.fill.fore_color.rgb = COLOR_ACCENT
    banner_line.line.fill.background()
    
    tx_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.18), Inches(9.2), Inches(0.75))
    tf_t = tx_title.text_frame
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.bold = True
    p_t.font.size = Pt(23)
    p_t.font.color.rgb = COLOR_TEXT_LIGHT
    
    tx_tag = slide.shapes.add_textbox(Inches(9.0), Inches(0.18), Inches(3.6), Inches(0.75))
    tf_tag = tx_tag.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = topic_tag
    p_tag.alignment = PP_ALIGN.RIGHT
    p_tag.font.bold = True
    p_tag.font.size = Pt(13)
    p_tag.font.color.rgb = COLOR_ACCENT
    
    # Cards layout
    if col3_title:
        card_w = Inches(3.7)
        card_gap = Inches(0.3)
        cols = [(col1_title, col1_items), (col2_title, col2_items), (col3_title, col3_items)]
    else:
        card_w = Inches(5.7)
        card_gap = Inches(0.4)
        cols = [(col1_title, col1_items), (col2_title, col2_items)]
        
    start_x = Inches(0.8)
    for idx, (c_title, c_items) in enumerate(cols):
        cx = start_x + idx * (card_w + card_gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.4), card_w, Inches(5.3))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG if idx % 2 == 0 else COLOR_CARD_ACCENT_BG
        card.line.color.rgb = COLOR_SECONDARY if idx == 0 else COLOR_PRIMARY
        card.line.width = Pt(1.5)
        
        tx = slide.shapes.add_textbox(cx + Inches(0.2), Inches(1.55), card_w - Inches(0.4), Inches(5.0))
        tf = tx.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = c_title
        p0.font.bold = True
        p0.font.size = Pt(16)
        p0.font.color.rgb = COLOR_PRIMARY if idx % 2 == 0 else COLOR_SECONDARY
        p0.space_after = Pt(12)
        
        for item in c_items:
            p = tf.add_paragraph()
            if isinstance(item, tuple):
                p.text = f"• {item[0]}: {item[1]}"
            else:
                p.text = f"• {item}"
            p.font.size = Pt(13)
            p.font.color.rgb = COLOR_TEXT_DARK
            p.space_after = Pt(6)
            
    # Footer
    footer = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.8), Inches(0.4))
    tf_f = footer.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = f"{COURSE_TAG} | {AUTHOR_TAG}"
    p_f.font.size = Pt(9.5)
    p_f.font.color.rgb = COLOR_TEXT_MUTED
    
    return slide

def init_prs():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs
